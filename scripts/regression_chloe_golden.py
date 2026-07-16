"""Golden regression runner for CHLOE exemplary input.

The script runs the package-level Excel helper with ``Exemplary_Inputs.xlsx`` and
compares the calculated values against a golden CSV generated from the trusted
baseline branch.

Usage:
    python scripts/regression_chloe_golden.py --update-golden
    python scripts/regression_chloe_golden.py
"""

from __future__ import annotations

import argparse
import csv
import math
import sys
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Any


REPO_ROOT = Path(__file__).resolve().parents[1]
SRC_ROOT = REPO_ROOT / "src"
if SRC_ROOT.exists() and str(SRC_ROOT) not in sys.path:
    sys.path.insert(0, str(SRC_ROOT))
DEFAULT_INPUT = REPO_ROOT / "Exemplary_Inputs.xlsx"
DEFAULT_GOLDEN = REPO_ROOT / "tests" / "golden" / "chloe_exemplary_outputs.csv"
DEFAULT_COMPARISON = REPO_ROOT / "regression_results" / "chloe_exemplary_comparison.xlsx"

PRINTED_RESULT_LABELS = {
    "heating_cooling_load_result": "total_heating_load_w_rounded",
    "heating_cooling_load_result2": "ventilation_losses_heating_w_rounded",
    "heating_cooling_load_result3": "solar_heat_gains_july_cooling_w_rounded",
    "heating_cooling_load_result4": "transmission_heat_gains_july_cooling_w_rounded",
    "heating_cooling_load_result5": "ventilation_heat_gains_july_cooling_w_rounded",
    "heating_cooling_load_result6": "transmission_losses_heating_w_rounded",
    "heating_cooling_load_result7": "total_cooling_load_w_rounded",
    "heating_cooling_load_result8": "solar_heat_gains_september_cooling_w_rounded",
    "heating_cooling_load_result9": "transmission_heat_gains_september_cooling_w_rounded",
    "heating_cooling_load_result10": "ventilation_heat_gains_september_cooling_w_rounded",
    "heating_cooling_load_result11": "total_cooling_load_july_w_rounded",
    "heating_cooling_load_result12": "total_cooling_load_september_w_rounded",
    "heating_cooling_load_result13": "internal_gains_cooling_w_rounded",
}


@dataclass(frozen=True)
class ResultRow:
    group: str
    key: str
    value: str
    value_type: str


@dataclass(frozen=True)
class Difference:
    group: str
    key: str
    golden_value: str
    current_value: str
    delta: str
    reason: str


def _stringify(value: Any) -> tuple[str, str]:
    if value is None:
        return "", "none"
    if isinstance(value, bool):
        return str(value), "bool"
    if isinstance(value, int) and not isinstance(value, bool):
        return str(value), "int"
    if isinstance(value, float):
        if math.isnan(value):
            return "nan", "float"
        if math.isinf(value):
            return "inf" if value > 0 else "-inf", "float"
        return repr(value), "float"
    return str(value), "str"


def _parse_numeric(value: str) -> float | None:
    if value == "":
        return None
    try:
        return float(value)
    except ValueError:
        return None


def _values_equal(golden: ResultRow, current: ResultRow, tolerance: float) -> tuple[bool, str]:
    if golden.value_type != current.value_type:
        return False, "type_changed"

    if golden.value_type in {"float", "int"}:
        golden_number = _parse_numeric(golden.value)
        current_number = _parse_numeric(current.value)
        if golden_number is None or current_number is None:
            return golden.value == current.value, "numeric_parse_fallback"
        if math.isnan(golden_number) and math.isnan(current_number):
            return True, "equal_nan"
        delta = abs(golden_number - current_number)
        return delta <= tolerance, f"delta={delta}"

    return golden.value == current.value, "exact"


def _run_simulation(input_path: Path) -> tuple[dict[str, Any], float]:
    if not input_path.exists():
        raise FileNotFoundError(f"input file not found: {input_path}")

    from chloe.excel_io import run_chloe

    start = time.perf_counter()
    result = run_chloe(str(input_path))
    elapsed = time.perf_counter() - start

    namespace = {
        "input_values": result.get("input_values"),
        "calculator": result.get("calculator"),
    }

    rounded_outputs = result.get("rounded_outputs", {})
    if isinstance(rounded_outputs, dict):
        reverse_labels = {
            "total_heating_load_w_rounded": "total_heating_load",
            "ventilation_losses_heating_w_rounded": "ventilation_losses_heating",
            "solar_heat_gains_july_cooling_w_rounded": "solar_heat_gains_july_cooling",
            "transmission_heat_gains_july_cooling_w_rounded": "transmission_heat_gains_july_cooling",
            "ventilation_heat_gains_july_cooling_w_rounded": "ventilation_heat_gains_july_cooling",
            "transmission_losses_heating_w_rounded": "transmission_losses_heating",
            "total_cooling_load_w_rounded": "total_cooling_load",
            "solar_heat_gains_september_cooling_w_rounded": "solar_heat_gains_september_cooling",
            "transmission_heat_gains_september_cooling_w_rounded": "transmission_heat_gains_september_cooling",
            "ventilation_heat_gains_september_cooling_w_rounded": "ventilation_heat_gains_september_cooling",
            "total_cooling_load_july_w_rounded": "total_cooling_load_july",
            "total_cooling_load_september_w_rounded": "total_cooling_load_september",
            "internal_gains_cooling_w_rounded": "internal_gains_cooling",
        }
        for variable_name, output_key in PRINTED_RESULT_LABELS.items():
            result_key = reverse_labels[output_key]
            if result_key in rounded_outputs:
                namespace[variable_name] = rounded_outputs[result_key]
    return namespace, elapsed


def _collect_results(namespace: dict[str, Any]) -> list[ResultRow]:
    rows: list[ResultRow] = []

    input_values = namespace.get("input_values")
    if isinstance(input_values, dict):
        for key in sorted(input_values):
            value, value_type = _stringify(input_values[key])
            rows.append(ResultRow("input", key, value, value_type))

    calculator = namespace.get("calculator")
    calculator_values = getattr(calculator, "__dict__", {})
    if isinstance(calculator_values, dict):
        for key in sorted(calculator_values):
            value, value_type = _stringify(calculator_values[key])
            rows.append(ResultRow("calculator", key, value, value_type))

    for variable_name, output_key in PRINTED_RESULT_LABELS.items():
        if variable_name in namespace:
            value, value_type = _stringify(namespace[variable_name])
            rows.append(ResultRow("rounded_output", output_key, value, value_type))

    if not rows:
        raise RuntimeError("No CHLOE results collected from package regression output.")
    return rows


def _write_csv(path: Path, rows: list[ResultRow]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as file:
        writer = csv.DictWriter(file, fieldnames=["group", "key", "value", "value_type"])
        writer.writeheader()
        for row in rows:
            writer.writerow(
                {
                    "group": row.group,
                    "key": row.key,
                    "value": row.value,
                    "value_type": row.value_type,
                }
            )


def _read_csv(path: Path) -> list[ResultRow]:
    with path.open("r", newline="", encoding="utf-8") as file:
        reader = csv.DictReader(file)
        return [
            ResultRow(
                group=row["group"],
                key=row["key"],
                value=row["value"],
                value_type=row["value_type"],
            )
            for row in reader
        ]


def _compare(golden_rows: list[ResultRow], current_rows: list[ResultRow], tolerance: float) -> list[Difference]:
    golden_by_key = {(row.group, row.key): row for row in golden_rows}
    current_by_key = {(row.group, row.key): row for row in current_rows}
    differences: list[Difference] = []

    for key in sorted(golden_by_key.keys() | current_by_key.keys()):
        golden = golden_by_key.get(key)
        current = current_by_key.get(key)
        group, result_key = key
        if golden is None and current is not None:
            differences.append(Difference(group, result_key, "", current.value, "", "added"))
            continue
        if current is None and golden is not None:
            differences.append(Difference(group, result_key, golden.value, "", "", "missing"))
            continue
        assert golden is not None and current is not None
        equal, reason = _values_equal(golden, current, tolerance)
        if not equal:
            delta = ""
            golden_number = _parse_numeric(golden.value)
            current_number = _parse_numeric(current.value)
            if golden_number is not None and current_number is not None:
                delta = repr(current_number - golden_number)
            differences.append(
                Difference(group, result_key, golden.value, current.value, delta, reason)
            )
    return differences


def _write_comparison(path: Path, golden_rows: list[ResultRow], current_rows: list[ResultRow], differences: list[Difference]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        import openpyxl
    except ImportError:
        fallback = path.with_suffix(".csv")
        with fallback.open("w", newline="", encoding="utf-8") as file:
            writer = csv.DictWriter(
                file,
                fieldnames=["group", "key", "golden_value", "current_value", "delta", "reason"],
            )
            writer.writeheader()
            for diff in differences:
                writer.writerow(diff.__dict__)
        return

    workbook = openpyxl.Workbook()
    summary = workbook.active
    summary.title = "summary"
    summary.append(["metric", "value"])
    summary.append(["golden_rows", len(golden_rows)])
    summary.append(["current_rows", len(current_rows)])
    summary.append(["differences", len(differences)])

    diff_sheet = workbook.create_sheet("differences")
    diff_sheet.append(["group", "key", "golden_value", "current_value", "delta", "reason"])
    for diff in differences:
        diff_sheet.append([
            diff.group,
            diff.key,
            diff.golden_value,
            diff.current_value,
            diff.delta,
            diff.reason,
        ])

    current_sheet = workbook.create_sheet("current")
    current_sheet.append(["group", "key", "value", "value_type"])
    for row in current_rows:
        current_sheet.append([row.group, row.key, row.value, row.value_type])

    workbook.save(path)


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Run CHLOE golden regression.")
    parser.add_argument("--update-golden", action="store_true", help="Write the current output as golden reference.")
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT, help="Path to Exemplary_Inputs.xlsx.")
    parser.add_argument("--golden", type=Path, default=DEFAULT_GOLDEN, help="Golden CSV path.")
    parser.add_argument("--comparison", type=Path, default=DEFAULT_COMPARISON, help="Comparison report path.")
    parser.add_argument("--tolerance", type=float, default=1e-9, help="Numeric comparison tolerance.")
    return parser.parse_args()


def main() -> int:
    args = _parse_args()
    namespace, elapsed_s = _run_simulation(args.input)
    current_rows = _collect_results(namespace)


    if args.update_golden:
        _write_csv(args.golden, current_rows)
        print(f"golden_updated={args.golden}")
        print(f"rows={len(current_rows)} elapsed_s={elapsed_s:.6f}")
        return 0

    if not args.golden.exists():
        print(f"golden_missing={args.golden}", file=sys.stderr)
        print("Run with --update-golden on the main branch first.", file=sys.stderr)
        return 2

    golden_rows = _read_csv(args.golden)
    differences = _compare(golden_rows, current_rows, args.tolerance)
    _write_comparison(args.comparison, golden_rows, current_rows, differences)

    print(f"rows={len(current_rows)} differences={len(differences)} elapsed_s={elapsed_s:.6f}")
    print(f"comparison_file={args.comparison}")
    if differences:
        for diff in differences[:10]:
            print(
                "difference "
                f"group={diff.group} key={diff.key} "
                f"golden={diff.golden_value} current={diff.current_value} "
                f"delta={diff.delta} reason={diff.reason}"
            )
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())



